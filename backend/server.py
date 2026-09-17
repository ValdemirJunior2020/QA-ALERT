from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from fastapi import File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import backend.app as app_module
from backend.app import app, current_user_from_request, db
from backend.booking_rules import (
    corrected_ingest_json_file,
    ensure_booking_schema,
    read_automation_settings,
    write_automation_settings,
)
from backend.knowledge_loader import KNOWLEDGE, source_folder, source_status
from backend.transcription_worker import STATUS_FILE, TRANSCRIPTS, ensure_schema

ROOT = Path(__file__).resolve().parent.parent
AUTOMATION_HTML = ROOT / "frontend" / "automation.html"
QA_STATUS_FILE = ROOT / "data" / "qa-worker-status.json"
EXPORTS = ROOT / "data" / "exports"
EXPORTS.mkdir(parents=True, exist_ok=True)

ensure_schema()
ensure_booking_schema()

# Permanent corrected ingest behavior used by the existing inbox watcher.
# scan_inbox_once() in backend.app resolves ingest_json_file dynamically, so
# replacing the module function here also repairs all future automatic scans.
app_module.ingest_json_file = corrected_ingest_json_file


def _worker_file_status():
    try:
        if STATUS_FILE.exists():
            return json.loads(STATUS_FILE.read_text(encoding="utf-8"))
    except Exception as exc:
        return {"state": "status_read_error", "error": str(exc)}
    return {"state": "not_started", "message": "Whisper worker has not written a status yet."}


def _qa_worker_status():
    try:
        if QA_STATUS_FILE.exists():
            return json.loads(QA_STATUS_FILE.read_text(encoding="utf-8"))
    except Exception as exc:
        return {"state": "status_read_error", "error": str(exc)}
    return {"state": "not_started", "message": "QA worker has not written a status yet."}


def _cell(row, key, default=None):
    try:
        value = row[key]
        return default if value is None else value
    except Exception:
        return default


def _transcription_status(row) -> str:
    if str(_cell(row, "transcript_text", "")).strip():
        return "Completed"
    if _cell(row, "transcription_error"):
        return "Error"
    if _cell(row, "status") == "transcribing":
        return "Listening"
    return "Waiting"


def _qa_status(row) -> str:
    status = str(_cell(row, "status", ""))
    if _cell(row, "qa_result_json"):
        if status == "needs_attention":
            return "Completed - Needs Attention"
        if status == "completed_no_booking":
            return "Behavior QA Completed"
        return "Completed"
    if status in {"qa_running"}:
        return "Analyzing"
    if str(_cell(row, "transcript_text", "")).strip():
        return "Waiting for QA"
    return "Waiting for transcript"


def _booking_state(row) -> str:
    if not int(_cell(row, "booking_found", 0) or 0):
        return "Not found"
    count = int(_cell(row, "booking_count", 0) or 0)
    state = str(_cell(row, "booking_match_status", "") or "")
    if count > 1 and state in {"pending", "ambiguous", "low_confidence"}:
        return "Multiple - needs one match"
    if state == "matched" or count == 1:
        return "Matched"
    return state or "Found"


def _build_full_export() -> Path:
    with db() as c:
        rows = c.execute("SELECT * FROM cases ORDER BY created_at DESC, id DESC").fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "QA Cases"

    headers = [
        "Call ID", "Booking", "Agent", "Call Center", "Caller Number", "Duration Sec",
        "Booking Found", "Booking Count", "Booking Match", "Match Confidence",
        "Transcription", "Transcript Language", "Transcript", "Audio Deleted",
        "QA Status", "Guest Request", "Agent Action", "Score", "Severity", "QA Finding",
        "QA Process", "QA Source", "Missing Notes", "Evidence Start", "Evidence End",
        "Evidence Excerpt", "Slack Alert Sent", "Created At", "Updated At",
    ]
    ws.append(headers)

    for r in rows:
        transcript = str(_cell(r, "transcript_text", "") or "")
        # Excel cells have a 32,767-character limit.
        if len(transcript) > 32000:
            transcript = transcript[:32000] + "\n[Transcript truncated in Excel; full TXT/JSON remains saved locally.]"
        ws.append([
            _cell(r, "call_id", "N/A"),
            _cell(r, "itinerary", "N/A") or "N/A",
            _cell(r, "agent", "N/A"),
            _cell(r, "call_center", "N/A"),
            _cell(r, "caller_number", "N/A"),
            _cell(r, "duration_seconds", 0),
            "Yes" if int(_cell(r, "booking_found", 0) or 0) else "No",
            _cell(r, "booking_count", 0),
            _booking_state(r),
            _cell(r, "booking_match_confidence"),
            _transcription_status(r),
            _cell(r, "transcript_language"),
            transcript,
            "Yes" if _cell(r, "audio_deleted_at") else "No",
            _qa_status(r),
            _cell(r, "guest_request"),
            _cell(r, "agent_action"),
            _cell(r, "score"),
            _cell(r, "severity", "info"),
            _cell(r, "finding"),
            _cell(r, "qa_process"),
            _cell(r, "qa_matrix_source"),
            "Yes" if int(_cell(r, "missing_notes", 0) or 0) else "No",
            _cell(r, "evidence_start_sec"),
            _cell(r, "evidence_end_sec"),
            _cell(r, "evidence_excerpt"),
            _cell(r, "slack_alert_sent_at"),
            _cell(r, "created_at"),
            _cell(r, "updated_at"),
        ])

    navy = "17213D"
    light = "F4F7FB"
    green = "DCFCE7"
    amber = "FEF3C7"
    red = "FEE2E2"
    purple = "EDE9FE"
    border_color = "D8DEEA"
    thin = Side(style="thin", color=border_color)

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row_idx in range(2, ws.max_row + 1):
        if row_idx % 2 == 0:
            for cell in ws[row_idx]:
                cell.fill = PatternFill("solid", fgColor=light)
        for cell in ws[row_idx]:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

        transcription_cell = ws.cell(row=row_idx, column=headers.index("Transcription") + 1)
        qa_cell = ws.cell(row=row_idx, column=headers.index("QA Status") + 1)
        severity_cell = ws.cell(row=row_idx, column=headers.index("Severity") + 1)
        source_cell = ws.cell(row=row_idx, column=headers.index("QA Source") + 1)

        transcription_cell.fill = PatternFill("solid", fgColor=green if transcription_cell.value == "Completed" else (red if transcription_cell.value == "Error" else amber))
        if str(qa_cell.value).startswith("Completed") or qa_cell.value == "Behavior QA Completed":
            qa_cell.fill = PatternFill("solid", fgColor=green)
        elif qa_cell.value == "Analyzing":
            qa_cell.fill = PatternFill("solid", fgColor=amber)
        else:
            qa_cell.fill = PatternFill("solid", fgColor=amber)
        if str(severity_cell.value).lower() == "critical":
            severity_cell.fill = PatternFill("solid", fgColor=red)
            severity_cell.font = Font(bold=True)
        elif str(severity_cell.value).lower() == "warning":
            severity_cell.fill = PatternFill("solid", fgColor=amber)
        if source_cell.value == "Matrix update emails sent":
            source_cell.fill = PatternFill("solid", fgColor=purple)

    left_headers = {"Transcript", "Guest Request", "Agent Action", "QA Finding", "QA Process", "Evidence Excerpt"}
    for header in left_headers:
        col = headers.index(header) + 1
        for row_idx in range(2, ws.max_row + 1):
            ws.cell(row=row_idx, column=col).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    widths = {
        "Call ID": 34, "Booking": 16, "Agent": 28, "Call Center": 14, "Caller Number": 20,
        "Duration Sec": 12, "Booking Found": 13, "Booking Count": 13, "Booking Match": 23,
        "Match Confidence": 16, "Transcription": 16, "Transcript Language": 16, "Transcript": 48,
        "Audio Deleted": 14, "QA Status": 24, "Guest Request": 32, "Agent Action": 32,
        "Score": 10, "Severity": 12, "QA Finding": 38, "QA Process": 34, "QA Source": 26,
        "Missing Notes": 14, "Evidence Start": 14, "Evidence End": 14, "Evidence Excerpt": 38,
        "Slack Alert Sent": 23, "Created At": 24, "Updated At": 24,
    }
    for idx, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(header, 16)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30

    details = wb.create_sheet("Details")
    detail_headers = [
        "Call ID", "Source JSON", "Audio Path", "Transcript TXT", "Transcript JSON",
        "Matched Booking JSON", "QA Result JSON", "Evidence File", "Transcription Error",
        "Transcription Started", "Transcription Completed", "Audio Deleted At",
    ]
    details.append(detail_headers)
    for r in rows:
        details.append([
            _cell(r, "call_id"), _cell(r, "source_json_path"), _cell(r, "audio_path"),
            _cell(r, "transcript_path"), _cell(r, "transcript_json_path"), _cell(r, "matched_booking_path"),
            _cell(r, "qa_result_json"), _cell(r, "evidence_path"), _cell(r, "transcription_error"),
            _cell(r, "transcription_started_at"), _cell(r, "transcription_completed_at"), _cell(r, "audio_deleted_at"),
        ])
    for cell in details[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in details.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    details.freeze_panes = "A2"
    details.auto_filter.ref = details.dimensions
    for idx in range(1, len(detail_headers) + 1):
        details.column_dimensions[get_column_letter(idx)].width = 28 if idx > 1 else 34

    out = EXPORTS / f"QA-ALERT-ALL-CASES-{datetime.now().strftime('%Y-%m-%d_%H%M%S')}.xlsx"
    wb.save(out)
    return out


# The older app.py export was created before transcription/QA columns existed.
# Intercept the same Download Excel URL so the button always exports the complete live database.
@app.middleware("http")
async def complete_excel_export(request: Request, call_next):
    if request.url.path != "/api/export.xlsx":
        return await call_next(request)
    if not current_user_from_request(request):
        return await call_next(request)
    out = _build_full_export()
    return FileResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=out.name,
    )


@app.get("/automation")
def automation_page():
    if not AUTOMATION_HTML.exists():
        raise HTTPException(404, "Automation settings page missing")
    return FileResponse(AUTOMATION_HTML)


@app.get("/api/automation-settings")
def get_automation_settings():
    return read_automation_settings()


@app.put("/api/automation-settings")
async def put_automation_settings(payload: dict):
    return write_automation_settings(payload)


@app.get("/api/knowledge/status")
def knowledge_status():
    return source_status()


@app.post("/api/knowledge/upload")
async def knowledge_upload(source_type: str = Form(...), file: UploadFile = File(...)):
    filename = Path(file.filename or "").name
    ext = Path(filename).suffix.lower()
    if ext not in {".xlsx", ".xlsm", ".docx", ".txt", ".md", ".json"}:
        raise HTTPException(400, "Use XLSX, XLSM, DOCX, TXT, MD, or JSON")
    if not filename:
        raise HTTPException(400, "File name is missing")
    try:
        folder = source_folder(source_type)
    except ValueError:
        raise HTTPException(400, "Choose QA Form / Rubric, Original Matrix, or Matrix update emails sent")

    content = await file.read()
    if len(content) > 50 * 1024 * 1024:
        raise HTTPException(400, "Knowledge file is larger than 50 MB")

    for old in folder.iterdir():
        if old.is_file():
            try:
                old.unlink()
            except Exception:
                pass

    target = folder / filename
    target.write_bytes(content)
    return {"ok": True, "source_type": source_type, "saved": filename, "status": source_status()}


@app.get("/api/workers")
def worker_status():
    with db() as c:
        queued = c.execute("SELECT COUNT(*) n FROM cases WHERE status='queued'").fetchone()["n"]
        transcribing = c.execute("SELECT COUNT(*) n FROM cases WHERE status='transcribing'").fetchone()["n"]
        completed = c.execute("SELECT COUNT(*) n FROM cases WHERE status='transcribed_audio_deleted'").fetchone()["n"]
        cleanup_errors = c.execute("SELECT COUNT(*) n FROM cases WHERE status='transcribed_audio_cleanup_error'").fetchone()["n"]
        errors = c.execute("SELECT COUNT(*) n FROM cases WHERE status='transcription_error'").fetchone()["n"]
        qa_running = c.execute("SELECT COUNT(*) n FROM cases WHERE status='qa_running'").fetchone()["n"]
        qa_errors = c.execute("SELECT COUNT(*) n FROM cases WHERE status='qa_error'").fetchone()["n"]
        completed_qa = c.execute("SELECT COUNT(*) n FROM cases WHERE status='completed'").fetchone()["n"]
        last = c.execute(
            """
            SELECT id,call_id,agent,call_center,status,transcript_language,
                   transcription_started_at,transcription_completed_at,audio_deleted_at,
                   transcription_error,transcript_path,itinerary,booking_count,booking_match_status,
                   booking_match_confidence,qa_process,qa_matrix_source,slack_alert_sent_at,updated_at
            FROM cases
            WHERE transcription_started_at IS NOT NULL OR transcript_text IS NOT NULL
            ORDER BY COALESCE(transcription_completed_at,transcription_started_at,updated_at) DESC
            LIMIT 1
            """
        ).fetchone()
    return {
        "whisper": _worker_file_status(),
        "qa": _qa_worker_status(),
        "knowledge": source_status(),
        "counts": {
            "queued": queued,
            "transcribing": transcribing,
            "completed_audio_deleted": completed,
            "cleanup_errors": cleanup_errors,
            "transcription_errors": errors,
            "qa_running": qa_running,
            "qa_errors": qa_errors,
            "completed_qa": completed_qa,
        },
        "last_case": dict(last) if last else None,
    }


@app.get("/api/cases/recent")
def recent_cases(limit: int = 50):
    limit = max(1, min(200, int(limit)))
    with db() as c:
        rows = c.execute(
            """
            SELECT id,call_id,itinerary,agent,call_center,duration_seconds,status,severity,
                   missing_notes,booking_found,booking_count,booking_match_status,booking_match_confidence,
                   transcript_language,transcription_started_at,transcription_completed_at,audio_deleted_at,
                   transcription_error,qa_process,qa_matrix_source,slack_alert_sent_at,created_at,updated_at
            FROM cases
            ORDER BY created_at DESC
            LIMIT ?
            """,
            (limit,),
        ).fetchall()
    return {"cases": [dict(r) for r in rows]}


@app.get("/api/cases/{case_id}/bookings")
def case_bookings(case_id: int):
    with db() as c:
        case = c.execute(
            "SELECT id,call_id,itinerary,booking_found,booking_count,booking_match_status,booking_match_confidence FROM cases WHERE id=?",
            (case_id,),
        ).fetchone()
        if not case:
            raise HTTPException(404, "Case not found")
        rows = c.execute(
            "SELECT id,itinerary,fit_id,booking_json_path,documentation_found,is_matched,match_confidence FROM booking_candidates WHERE case_id=? ORDER BY id",
            (case_id,),
        ).fetchall()
    return {"case": dict(case), "bookings": [dict(r) for r in rows]}


@app.get("/api/cases/{case_id}/transcript")
def case_transcript(case_id: int):
    with db() as c:
        row = c.execute(
            """
            SELECT id,call_id,agent,call_center,status,transcript_text,transcript_language,
                   transcript_path,transcript_json_path,transcription_completed_at,audio_deleted_at,
                   transcription_error,itinerary,booking_count,booking_match_status,qa_process,qa_matrix_source
            FROM cases WHERE id=?
            """,
            (case_id,),
        ).fetchone()
    if not row:
        raise HTTPException(404, "Case not found")
    return dict(row)


@app.post("/api/cases/{case_id}/transcription/retry")
def retry_transcription(case_id: int):
    with db() as c:
        row = c.execute("SELECT id,status,transcript_text FROM cases WHERE id=?", (case_id,)).fetchone()
        if not row:
            raise HTTPException(404, "Case not found")
        if row["transcript_text"]:
            raise HTTPException(409, "This case already has a saved transcript")
        c.execute(
            "UPDATE cases SET status='queued',transcription_error=NULL,updated_at=datetime('now') WHERE id=?",
            (case_id,),
        )
    return {"ok": True, "case_id": case_id, "status": "queued"}
