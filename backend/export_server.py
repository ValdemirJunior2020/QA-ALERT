from __future__ import annotations

from datetime import datetime
from pathlib import Path

from fastapi import Request
from fastapi.responses import FileResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.server import app, current_user_from_request, db

ROOT = Path(__file__).resolve().parent.parent
EXPORTS = ROOT / "data" / "exports"
EXPORTS.mkdir(parents=True, exist_ok=True)


def _v(row, key, default=None):
    try:
        value = row[key]
        return default if value is None else value
    except Exception:
        return default


def _transcription_state(row):
    if str(_v(row, "transcript_text", "") or "").strip():
        return "Completed"
    if _v(row, "status") == "transcribing":
        return "Listening"
    if _v(row, "transcription_error"):
        return "Error"
    return "Waiting"


def _qa_state(row):
    status = str(_v(row, "status", "") or "")
    if _v(row, "qa_result_json"):
        if status == "needs_attention":
            return "Completed - Needs Attention"
        if status == "completed_no_booking":
            return "Behavior QA Completed"
        return "Completed"
    if status in {"qa_running", "booking_matching", "matrix_check", "ollama_qa"}:
        return "Analyzing"
    if str(_v(row, "transcript_text", "") or "").strip():
        return "Waiting for QA"
    return "Waiting for transcript"


def _booking_state(row):
    if not int(_v(row, "booking_found", 0) or 0):
        return "Not found"
    count = int(_v(row, "booking_count", 0) or 0)
    state = str(_v(row, "booking_match_status", "") or "")
    if count > 1 and state in {"pending", "ambiguous", "low_confidence"}:
        return "Multiple - needs one match"
    if state == "matched" or count == 1:
        return "Matched"
    return state or "Found"


def build_export() -> tuple[Path, int]:
    with db() as c:
        count = c.execute("SELECT COUNT(*) n FROM cases").fetchone()["n"]
        rows = c.execute("SELECT * FROM cases ORDER BY created_at DESC, id DESC").fetchall()

    # The count and actual selected rows must agree. If they do not, fail loudly
    # instead of silently producing a partial workbook.
    if int(count) != len(rows):
        raise RuntimeError(f"Export row mismatch: database count={count}, selected rows={len(rows)}")

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["QA ALERT Export"])
    summary.append(["Database cases", int(count)])
    summary.append(["Exported case rows", len(rows)])
    summary.append(["Generated at", datetime.now().isoformat(timespec="seconds")])
    summary.append(["Verification", "MATCH" if int(count) == len(rows) else "MISMATCH"])
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 28
    summary["A1"].font = Font(bold=True, size=14)
    for cell in summary[1]:
        cell.fill = PatternFill("solid", fgColor="17213D")
        cell.font = Font(color="FFFFFF", bold=True)

    ws = wb.create_sheet("QA Cases")
    headers = [
        "Call ID", "Booking", "Agent", "Call Center", "Caller Number", "Duration Sec",
        "Booking Found", "Booking Count", "Booking Match", "Match Confidence",
        "Transcription", "Transcript Language", "Transcript", "Audio Deleted",
        "QA Status", "Guest Request", "Agent Action", "Score", "Severity", "QA Finding",
        "QA Process", "QA Source", "Missing Notes", "Evidence Start", "Evidence End",
        "Evidence Excerpt", "Slack Alert Sent", "Created At", "Updated At",
    ]
    ws.append(headers)

    for r in rows:
        transcript = str(_v(r, "transcript_text", "") or "")
        if len(transcript) > 32000:
            transcript = transcript[:32000] + "\n[Transcript truncated in Excel; full transcript remains saved locally.]"
        ws.append([
            _v(r, "call_id", "N/A"),
            _v(r, "itinerary", "N/A") or "N/A",
            _v(r, "agent", "N/A"),
            _v(r, "call_center", "N/A"),
            _v(r, "caller_number", "N/A"),
            _v(r, "duration_seconds", 0),
            "Yes" if int(_v(r, "booking_found", 0) or 0) else "No",
            _v(r, "booking_count", 0),
            _booking_state(r),
            _v(r, "booking_match_confidence"),
            _transcription_state(r),
            _v(r, "transcript_language"),
            transcript,
            "Yes" if _v(r, "audio_deleted_at") else "No",
            _qa_state(r),
            _v(r, "guest_request"),
            _v(r, "agent_action"),
            _v(r, "score"),
            _v(r, "severity", "info"),
            _v(r, "finding"),
            _v(r, "qa_process"),
            _v(r, "qa_matrix_source"),
            "Yes" if int(_v(r, "missing_notes", 0) or 0) else "No",
            _v(r, "evidence_start_sec"),
            _v(r, "evidence_end_sec"),
            _v(r, "evidence_excerpt"),
            _v(r, "slack_alert_sent_at"),
            _v(r, "created_at"),
            _v(r, "updated_at"),
        ])

    navy = "17213D"
    light = "F4F7FB"
    green = "DCFCE7"
    amber = "FEF3C7"
    red = "FEE2E2"
    purple = "EDE9FE"
    thin = Side(style="thin", color="D8DEEA")

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row_idx in range(2, ws.max_row + 1):
        for cell in ws[row_idx]:
            if row_idx % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=light)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

        t = ws.cell(row=row_idx, column=headers.index("Transcription") + 1)
        q = ws.cell(row=row_idx, column=headers.index("QA Status") + 1)
        sev = ws.cell(row=row_idx, column=headers.index("Severity") + 1)
        src = ws.cell(row=row_idx, column=headers.index("QA Source") + 1)
        t.fill = PatternFill("solid", fgColor=green if t.value == "Completed" else (red if t.value == "Error" else amber))
        q.fill = PatternFill("solid", fgColor=green if str(q.value).startswith("Completed") or q.value == "Behavior QA Completed" else amber)
        if str(sev.value).lower() == "critical":
            sev.fill = PatternFill("solid", fgColor=red)
            sev.font = Font(bold=True)
        elif str(sev.value).lower() == "warning":
            sev.fill = PatternFill("solid", fgColor=amber)
        if src.value == "Matrix update emails sent":
            src.fill = PatternFill("solid", fgColor=purple)

    for header in {"Transcript", "Guest Request", "Agent Action", "QA Finding", "QA Process", "Evidence Excerpt"}:
        col = headers.index(header) + 1
        for row_idx in range(2, ws.max_row + 1):
            ws.cell(row=row_idx, column=col).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    widths = {
        "Call ID": 34, "Booking": 16, "Agent": 28, "Call Center": 14, "Caller Number": 20,
        "Duration Sec": 12, "Booking Found": 13, "Booking Count": 13, "Booking Match": 23,
        "Match Confidence": 16, "Transcription": 16, "Transcript Language": 16, "Transcript": 48,
        "Audio Deleted": 14, "QA Status": 24, "Guest Request": 32, "Agent Action": 32,
        "Score": 10, "Severity": 12, "QA Finding": 38, "QA Process": 34, "QA Source": 26,
        "Missing Notes": 14, "Evidence Start": 14, "Evidence End": 14, "Evidence Excerpt": 38,
        "Slack Alert Sent": 23, "Created At": 24, "Updated At": 24,
    }
    for idx, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(header, 16)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30

    details = wb.create_sheet("Details")
    detail_headers = [
        "Call ID", "Status", "Source JSON", "Audio Path", "Transcript TXT", "Transcript JSON",
        "Matched Booking JSON", "QA Result JSON", "Evidence File", "Transcription Error",
        "Transcription Started", "Transcription Completed", "Audio Deleted At",
    ]
    details.append(detail_headers)
    for r in rows:
        details.append([
            _v(r, "call_id"), _v(r, "status"), _v(r, "source_json_path"), _v(r, "audio_path"),
            _v(r, "transcript_path"), _v(r, "transcript_json_path"), _v(r, "matched_booking_path"),
            _v(r, "qa_result_json"), _v(r, "evidence_path"), _v(r, "transcription_error"),
            _v(r, "transcription_started_at"), _v(r, "transcription_completed_at"), _v(r, "audio_deleted_at"),
        ])
    for cell in details[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    details.freeze_panes = "A2"
    details.auto_filter.ref = details.dimensions
    for idx in range(1, len(detail_headers) + 1):
        details.column_dimensions[get_column_letter(idx)].width = 28 if idx > 2 else 34

    out = EXPORTS / f"QA-ALERT-{int(count)}-CASES-{datetime.now().strftime('%Y-%m-%d_%H%M%S')}.xlsx"
    wb.save(out)
    return out, int(count)


@app.middleware("http")
async def guaranteed_full_export(request: Request, call_next):
    if request.url.path not in {"/api/export.xlsx", "/api/export-all.xlsx"}:
        return await call_next(request)
    if not current_user_from_request(request):
        return await call_next(request)
    out, count = build_export()
    return FileResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=out.name,
        headers={
            "Cache-Control": "no-store, no-cache, must-revalidate, max-age=0",
            "Pragma": "no-cache",
            "X-QA-Alert-Case-Count": str(count),
        },
    )


@app.get("/api/export-check")
def export_check():
    with db() as c:
        total = c.execute("SELECT COUNT(*) n FROM cases").fetchone()["n"]
        statuses = [dict(r) for r in c.execute("SELECT status,COUNT(*) count FROM cases GROUP BY status ORDER BY status").fetchall()]
    return {"database_cases": int(total), "statuses": statuses}
