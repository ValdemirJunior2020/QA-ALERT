from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
KNOWLEDGE = ROOT / "knowledge"
KNOWLEDGE.mkdir(parents=True, exist_ok=True)

TEXT_EXTS = {".txt", ".md", ".json"}
EXCEL_EXTS = {".xlsx", ".xlsm"}
WORD_EXTS = {".docx"}


def _read_text(path: Path) -> str:
    if path.suffix.lower() == ".json":
        try:
            obj = json.loads(path.read_text(encoding="utf-8-sig", errors="ignore"))
            return json.dumps(obj, ensure_ascii=False, indent=2)
        except Exception:
            pass
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_excel(path: Path) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    chunks: list[str] = []
    for ws in wb.worksheets:
        chunks.append(f"\n### SHEET: {ws.title}\n")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v).strip() for v in row if v not in (None, "")]
            if vals:
                chunks.append(" | ".join(vals))
    return "\n".join(chunks)


def _read_docx(path: Path) -> str:
    doc = Document(path)
    chunks: list[str] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            chunks.append(text)
    for table in doc.tables:
        for row in table.rows:
            vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if vals:
                chunks.append(" | ".join(vals))
    return "\n".join(chunks)


def source_label(path: Path) -> str:
    name = path.name.lower()
    if "email" in name or "issue" in name or "update" in name:
        return "Matrix update emails sent"
    if "original" in name and "matrix" in name:
        return "Original Matrix"
    if "qa" in name or "rubric" in name or "fly" in name:
        return "QA Form / Rubric"
    return path.stem


def load_knowledge_text(max_chars_per_file: int = 160000) -> str:
    chunks: list[str] = []
    for path in sorted(KNOWLEDGE.rglob("*")):
        if not path.is_file() or path.name.startswith("~$"):
            continue
        ext = path.suffix.lower()
        try:
            if ext in TEXT_EXTS:
                text = _read_text(path)
            elif ext in EXCEL_EXTS:
                text = _read_excel(path)
            elif ext in WORD_EXTS:
                text = _read_docx(path)
            else:
                continue
        except Exception as exc:
            chunks.append(f"\n--- {source_label(path)} (read error: {exc}) ---\n")
            continue
        if text.strip():
            chunks.append(f"\n--- SOURCE: {source_label(path)} ---\n{text[:max_chars_per_file]}")
    return "\n".join(chunks)


def source_status() -> dict:
    files = []
    for path in sorted(KNOWLEDGE.rglob("*")):
        if path.is_file() and path.suffix.lower() in (TEXT_EXTS | EXCEL_EXTS | WORD_EXTS):
            files.append({"name": path.name, "label": source_label(path), "size": path.stat().st_size})
    labels = {x["label"] for x in files}
    return {
        "folder": str(KNOWLEDGE),
        "files": files,
        "has_qa_form": "QA Form / Rubric" in labels,
        "has_original_matrix": "Original Matrix" in labels,
        "has_matrix_updates": "Matrix update emails sent" in labels,
        "ready": all(x in labels for x in ("QA Form / Rubric", "Original Matrix", "Matrix update emails sent")),
    }
